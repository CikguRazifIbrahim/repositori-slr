"""Bina data/scopus.json daripada Scopus Source Title List (Excel awam Elsevier).

Guna:  python3 tools/bina_scopus.py "Scopus-Title-List.xlsx" data/scopus.json
Perlu: pip install openpyxl

Helaian "Scopus Active (...)" dibaca; hanya jurnal berstatus Active diambil. Kuartil dianggar
daripada lajur "SJR" tahun terkini mengikut kategori ASJC (kaedah SCImago: kedudukan SJR dalam
setiap kategori, Q1 = 25 peratus teratas; kuartil terbaik merentas kategori disimpan).
Nilai 0 bermakna jurnal Scopus tanpa nilai SJR.
"""
import bisect
import collections
import json
import re
import sys
from datetime import date

import openpyxl


def issn8(x):
    x = re.sub(r"[^0-9Xx]", "", str(x or ""))
    return x.upper() if len(x) == 8 else ""


def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True)
    name = next((n for n in wb.sheetnames if n.lower().startswith("scopus active")), None)
    if not name:
        raise SystemExit("Helaian 'Scopus Active' tidak ditemui: " + ", ".join(wb.sheetnames))
    ws = wb[name]
    rows_iter = ws.iter_rows(values_only=True)
    hdr = [str(h or "").strip() for h in next(rows_iter)]
    col = {h.lower(): i for i, h in enumerate(hdr)}
    i_issn, i_eissn, i_active, i_type = col["issn"], col["eissn"], col["active"], col["source type"]
    i_sjr = next(i for h, i in col.items() if h.endswith(" sjr"))
    i_asjc = next(i for h, i in col.items() if "asjc" in h)
    rows = []
    for r in rows_iter:
        if not r or not r[0]:
            continue
        if str(r[i_active]).strip().lower() != "active" or str(r[i_type] or "").strip().lower() != "journal":
            continue
        issns = [x for x in (issn8(r[i_issn]), issn8(r[i_eissn])) if x]
        if not issns:
            continue
        try:
            sjr = float(r[i_sjr]) if r[i_sjr] not in (None, "", "None") else None
        except (TypeError, ValueError):
            sjr = None
        codes = [c.strip() for c in str(r[i_asjc] or "").split(";") if c.strip()]
        rows.append((issns, sjr, codes))
    bycat = collections.defaultdict(list)
    for issns, sjr, codes in rows:
        if sjr is not None:
            for c in codes:
                bycat[c].append(sjr)
    for c in bycat:
        bycat[c].sort()
    out = {}
    for issns, sjr, codes in rows:
        q = 0
        if sjr is not None:
            best = 5
            for c in codes:
                arr = bycat.get(c) or []
                if not arr:
                    continue
                higher = len(arr) - bisect.bisect_right(arr, sjr)
                best = min(best, int(higher * 4 / len(arr)) + 1)
            q = best if best < 5 else 0
        for i in issns:
            if i in out and out[i] and (not q or out[i] <= q):
                continue
            out[i] = q
    meta = {"sumber": "Scopus Source Title List (Elsevier, senarai awam), helaian " + name + "; kuartil dianggar daripada SJR mengikut kategori ASJC (kaedah SCImago)",
            "fail": src.split("/")[-1], "tarikh_bina": date.today().isoformat(), "jumlah_jurnal": len(rows)}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "q": out}, f, separators=(",", ":"))
    print("jurnal aktif:", len(rows), "| ISSN:", len(out), "| ->", dst)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
